import hashlib
import json
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from career_agent.documents.tailoring import DocumentError, validate_selection

SECTIONS = {
    "education": "Education",
    "skill": "Technical Skills",
    "experience": "Experience",
    "project": "Projects",
    "volunteer": "Volunteer Experience",
}


def compact(text):
    return re.sub(r"\s+", " ", text).strip()


def export_documents(profile, job, selection, output: Path):
    validate_selection(profile, selection)
    claims = {c.claim_id: c for c in profile.verified_claims}
    ordered = [claims[i] for i in selection.claim_ids]
    font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    if not font_path.exists():
        raise DocumentError("pdf_font_missing")
    if "CareerSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("CareerSans", str(font_path)))
    lines = [(profile.name, "name")]
    contact = " | ".join(
        v
        for v in [
            profile.location,
            profile.application_email or profile.email,
            profile.phone,
            profile.linkedin_url,
        ]
        if v
    )
    if contact:
        lines.append((contact, "body"))
    lines.append(("Target role: " + job.title, "body"))
    for category, heading in SECTIONS.items():
        selected = [c for c in ordered if c.category == category]
        if selected:
            lines.append((heading, "heading"))
            lines.extend((c.text, "body") for c in selected)
    languages = [
        f"{lang.language}: {lang.proficiency}"
        for lang in profile.languages
        if lang.verification.value == "VERIFIED"
    ]
    if languages:
        lines.extend([("Languages", "heading"), ("; ".join(languages), "body")])
    # Reject glyphs the actual font cannot render, rather than producing unreadable PDFs.
    cmap = pdfmetrics.getFont("CareerSans").face.charToGlyph
    if any(ord(ch) not in cmap for text, _ in lines for ch in text if not ch.isspace()):
        raise DocumentError("unsupported_pdf_glyphs")
    output.mkdir(mode=0o700, parents=True, exist_ok=False)
    document = Document()
    document.sections[0].top_margin = Inches(0.6)
    document.sections[0].bottom_margin = Inches(0.6)
    document.styles["Normal"].font.name = "DejaVu Sans"
    document.styles["Normal"].font.size = Pt(10)
    story = []
    for text, kind in lines:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = kind in {"name", "heading"}
        run.font.size = Pt(16 if kind == "name" else 11 if kind == "heading" else 10)
        style = ParagraphStyle(
            "resume",
            fontName="CareerSans",
            fontSize=16 if kind == "name" else 11 if kind == "heading" else 10,
            leading=14,
            spaceAfter=6,
        )
        story.extend([Paragraph(escape(text).replace("\n", "<br/>"), style), Spacer(1, 2)])
    document.save(output / "resume.docx")
    SimpleDocTemplate(
        str(output / "resume.pdf"), pagesize=letter, topMargin=42, bottomMargin=42
    ).build(story)
    cover = Document()
    cover.add_paragraph(f"Dear {job.company_name} hiring team,")
    cover.add_paragraph(
        f"I am applying for the {job.title} position. "
        "My verified experience includes the following:"
    )
    for claim in ordered[:5]:
        cover.add_paragraph(claim.text)
    cover.add_paragraph("Thank you for considering my application.")
    cover.add_paragraph(profile.name)
    cover.save(output / "cover_letter.docx")
    cover_story = [
        Paragraph(
            escape(p.text),
            ParagraphStyle("cover", fontName="CareerSans", fontSize=10, leading=14, spaceAfter=10),
        )
        for p in cover.paragraphs
    ]
    SimpleDocTemplate(str(output / "cover_letter.pdf"), pagesize=letter).build(cover_story)
    pdf_text = compact(
        " ".join(page.extract_text() or "" for page in PdfReader(output / "resume.pdf").pages)
    )
    docx_text = compact(" ".join(p.text for p in Document(output / "resume.docx").paragraphs))
    for claim in ordered:
        if compact(claim.text) not in pdf_text or compact(claim.text) not in docx_text:
            raise DocumentError("ats_text_validation_failed")
    manifest = {
        "claim_ids": selection.claim_ids,
        "profile_hash": hashlib.sha256(
            json.dumps(profile.model_dump(mode="json"), sort_keys=True).encode()
        ).hexdigest(),
        "job_content_hash": hashlib.sha256(
            json.dumps(
                job.model_dump(mode="json", exclude={"fetched_at", "published_at", "updated_at"}),
                sort_keys=True,
            ).encode()
        ).hexdigest(),
        "source_job_url": job.source_url,
        "files": {
            name: hashlib.sha256((output / name).read_bytes()).hexdigest()
            for name in ["resume.docx", "resume.pdf", "cover_letter.docx", "cover_letter.pdf"]
        },
    }
    (output / "manifest.json").write_text(json.dumps(manifest, sort_keys=True, indent=2))
    for file in output.iterdir():
        file.chmod(0o600)
    return manifest
